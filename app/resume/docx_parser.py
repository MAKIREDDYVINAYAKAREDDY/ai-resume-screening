import io
from pathlib import Path

from docx import Document


def _extract_document_text(
    document: Document,
) -> str:

    parts = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:

        for row in table.rows:

            cells = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    cells.append(text)

            if cells:
                parts.append(
                    " | ".join(cells)
                )

    return "\n".join(parts).strip()


def extract_text_from_docx(
    file_path: str,
) -> str:

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            file_path
        )

    document = Document(
        str(path)
    )

    return _extract_document_text(
        document
    )


def extract_text_from_docx_bytes(
    data: bytes,
) -> str:

    document = Document(
        io.BytesIO(data)
    )

    return _extract_document_text(
        document
    )
